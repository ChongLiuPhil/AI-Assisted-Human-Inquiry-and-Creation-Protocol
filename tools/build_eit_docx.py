from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "paper/submission/ethics-information-technology/MANUSCRIPT_BLINDED.md"
DEFAULT_OUTPUT = ROOT / "paper/submission/ethics-information-technology/_build/MANUSCRIPT_BLINDED.docx"


def set_cell_shading(cell, fill="EDEDED"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number_field(paragraph):
    run = paragraph.add_run()
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    instruction.text = " PAGE "
    field_separate = OxmlElement("w:fldChar")
    field_separate.set(qn("w:fldCharType"), "separate")
    cached_value = OxmlElement("w:t")
    cached_value.text = "1"
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.extend([field_begin, instruction, field_separate, cached_value, field_end])


def add_formatted_text(paragraph, text):
    pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|\`[^\`]+\`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(9)
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def is_table_separator(line):
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c or "") for c in cells)


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].lstrip().startswith("|"):
        rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
        i += 1
    if len(rows) >= 2 and is_table_separator(lines[start + 1]):
        rows.pop(1)
    return rows, i


def build(source: Path, output: Path):
    lines = source.read_text(encoding="utf-8").splitlines()
    doc = Document()

    sec = doc.sections[0]
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)

    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.title = ""
    doc.core_properties.subject = ""
    doc.core_properties.keywords = ""
    doc.core_properties.comments = ""
    doc.core_properties.category = ""

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.15

    for name, size in (("Title", 16), ("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
        styles[name].font.name = "Arial"
        styles[name].font.size = Pt(size)
        styles[name].font.bold = True

    styles["Heading 1"].paragraph_format.space_before = Pt(12)
    styles["Heading 1"].paragraph_format.space_after = Pt(6)
    styles["Heading 2"].paragraph_format.space_before = Pt(10)
    styles["Heading 2"].paragraph_format.space_after = Pt(4)
    styles["Heading 3"].paragraph_format.space_before = Pt(8)
    styles["Heading 3"].paragraph_format.space_after = Pt(3)

    in_code = False
    code_lines = []
    i = 0
    first_heading = True

    while i < len(lines):
        line = lines[i]

        if line.startswith("~~~"):
            if in_code:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.right_indent = Inches(0.3)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(6)
                run = p.add_run("\n".join(code_lines))
                run.font.name = "Courier New"
                run.font.size = Pt(9)
                in_code = False
                code_lines = []
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if not line.strip() or line.strip() == "---":
            i += 1
            continue

        if line.startswith("|") and i + 1 < len(lines) and lines[i + 1].startswith("|"):
            rows, nxt = parse_table(lines, i)
            cols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = "Table Grid"
            for r_idx, row in enumerate(rows):
                for c_idx, value in enumerate(row):
                    cell = table.cell(r_idx, c_idx)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    p = cell.paragraphs[0]
                    add_formatted_text(p, value)
                    for run in p.runs:
                        run.font.name = "Arial"
                        run.font.size = Pt(9)
                        if r_idx == 0:
                            run.bold = True
                    if r_idx == 0:
                        set_cell_shading(cell)
            if rows:
                set_repeat_table_header(table.rows[0])
            i = nxt
            continue

        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            add_formatted_text(p, line[2:].strip())
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            first_heading = False
            i += 1
            continue

        if line.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            add_formatted_text(p, line[3:].strip())
            i += 1
            continue

        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_formatted_text(p, line[4:].strip())
            i += 1
            continue

        if line.startswith("#### "):
            p = doc.add_paragraph(style="Heading 3")
            add_formatted_text(p, line[5:].strip())
            i += 1
            continue

        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.right_indent = Inches(0.2)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            add_formatted_text(p, line[2:].strip())
            for run in p.runs:
                run.italic = True
            i += 1
            continue

        if re.match(r"^- ", line):
            p = doc.add_paragraph(style="List Bullet")
            add_formatted_text(p, line[2:].strip())
            i += 1
            continue

        m = re.match(r"^(\d+)\.\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_formatted_text(p, m.group(2).strip())
            i += 1
            continue

        # Merge consecutive plain lines into one paragraph.
        buf = [line.strip()]
        j = i + 1
        while j < len(lines):
            nxt = lines[j]
            if (
                not nxt.strip()
                or nxt.strip() == "---"
                or nxt.startswith("#")
                or nxt.startswith("> ")
                or nxt.startswith("- ")
                or re.match(r"^\d+\.\s+", nxt)
                or nxt.startswith("|")
                or nxt.startswith("~~~")
            ):
                break
            buf.append(nxt.strip())
            j += 1
        p = doc.add_paragraph()
        add_formatted_text(p, " ".join(buf))
        i = j

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Blinded manuscript - venue-specific submission derivative | Page ")
    add_page_number_field(footer)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return output


if __name__ == "__main__":
    source = Path(sys.argv[1]) if len(sys.argv) > 1 else SOURCE
    output = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_OUTPUT
    out = build(source, output)
    print(out)
